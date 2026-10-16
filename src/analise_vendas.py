
from pathlib import Path
import pandas as pd



from openpyxl import load_workbook
from openpyxl.chart import Reference, BarChart, LineChart

# 1. Localiza a pasta do projeto
PASTA_PROJETO = Path(__file__).resolve().parent.parent

# 2. Define o caminho do Excel
ARQUIVO = PASTA_PROJETO / "dados" / "controle_vendas.xlsx"

PASTA_RESULTADOS = PASTA_PROJETO / "resultados"

ARQUIVO_SAIDA = (
    PASTA_RESULTADOS / "relatorio_analise_vendas.xlsx"
)

PASTA_RESULTADOS.mkdir(parents=True, exist_ok=True)

# 3. Importa a aba Listagem
df = pd.read_excel(
    ARQUIVO,
    sheet_name="Listagem",
    header=3
)

print("Dados Importados: ", df.shape)

# print(df)

print(df.head(3))
# print(df.tail(2))

# print(df.info())

# print(df.describe())

print(df.isnull().sum())

df["Valor Total"] = (df["Valor Unitário"]) * (df["Qtde"])

print(df)

#  4. ANÁLISE POR VENDEDOR
# # ==================================================

# resumo_vendedor = (
#     df.groupby("Vendedores", as_index=False)
#     ["Valor Total"]
#     .sum()
#     .sort_values("Valor Total", ascending=False)
# )

resumo_vendedor = (df.groupby("Vendedores", as_index=False)["Valor Total"].sum().sort_values("Valor Total", ascending=False))

print(resumo_vendedor)


# 4. Exibe os primeiros registros
# print(df.head())

# # 5. Exibe informações sobre as colunas
# print(df.info())


# # Quantidade de registros

# print("\n Total de registros:", len(df))

# # Vendedores existentes
# print("\nVendedores:")
# print(df["Vendedores"].value_counts())
# # Produtos existentes
# print("\nProdutos:")
# print(df["Produto"].value_counts())
# # Períodos existentes
# print("\nPeríodos:")
# print(df["Período"].value_counts())


# # Investiga registros repetidos
# duplicados = df.duplicated()

# print("\nQuantidade de linhas repetidas:")
# print(duplicados.sum())

# print("\nExemplos de linhas repetidas:")
# print(df[duplicados].head(10))



# # Mostra todas as linhas que fazem parte de grupos duplicados
# duplicados = df.duplicated(keep=False)

# print("\nTodos os registros envolvidos em duplicações:")
# print(
#     df.loc[duplicados].sort_values(
#         by=["Vendedores", "Período", "Produto"]
#     )
# )



# print(duplicados.head(10))
# print("Total de linhas marcadas:", duplicados.sum())

# ==================================================
# 5. ANÁLISE POR PRODUTO
# ==================================================

resumo_produto = (
    df.groupby("Produto", as_index=False)
    ["Valor Total"]
    .sum()
    .sort_values("Valor Total", ascending=False)
)


# ==================================================
# 6. ANÁLISE POR PERÍODO
# ==================================================

resumo_periodo = (
    df.groupby("Período", as_index=False)
    ["Valor Total"]
    .sum()
)

# Extrai o número da semana para ordenar corretamente
resumo_periodo["_ordem"] = (
    resumo_periodo["Período"]
    .str.extract(r"(\d+)")
    .astype(int)
)

resumo_periodo = (
    resumo_periodo
    .sort_values("_ordem")
    .drop(columns="_ordem")
)


# ==================================================
# 7. EXPORTAÇÃO PARA NOVAS ABAS DO EXCEL
# ==================================================

with pd.ExcelWriter(
    ARQUIVO_SAIDA,
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="Dados",
        index=False
    )

    resumo_vendedor.to_excel(
        writer,
        sheet_name="Por vendedor",
        index=False
    )

    resumo_produto.to_excel(
        writer,
        sheet_name="Por produto",
        index=False
    )

    resumo_periodo.to_excel(
        writer,
        sheet_name="Por período",
        index=False
    )


# ==================================================
# 8. INSERÇÃO DOS GRÁFICOS NO EXCEL
# ==================================================

wb = load_workbook(ARQUIVO_SAIDA)


def adicionar_grafico(
    nome_aba,
    titulo,
    tipo="bar"
):
    ws = wb[nome_aba]

    # Seleciona os valores da segunda coluna
    dados = Reference(
        ws,
        min_col=2,
        min_row=1,
        max_row=ws.max_row
    )

    # Seleciona os nomes das categorias
    categorias = Reference(
        ws,
        min_col=1,
        min_row=2,
        max_row=ws.max_row
    )

    # Escolhe o tipo de gráfico
    if tipo == "line":
        grafico = LineChart()
    else:
        grafico = BarChart()
        grafico.type = "bar"

    grafico.title = titulo
    grafico.y_axis.title = "Faturamento (R$)"
    grafico.x_axis.title = "Categoria"

    grafico.add_data(
        dados,
        titles_from_data=True
    )

    grafico.set_categories(categorias)

    grafico.height = 9
    grafico.width = 18

    # Insere o gráfico ao lado da tabela
    ws.add_chart(grafico, "D2")


# Gráfico de faturamento por vendedor
adicionar_grafico(
    "Por vendedor",
    "Faturamento por vendedor"
)

# Gráfico de faturamento por produto
adicionar_grafico(
    "Por produto",
    "Faturamento por produto"
)

# Gráfico de faturamento semanal
adicionar_grafico(
    "Por período",
    "Faturamento por semana",
    tipo="line"
)


# ==================================================
# 9. SALVAR O RELATÓRIO
# ==================================================

wb.save(ARQUIVO_SAIDA)

print("\nRelatório criado com sucesso!")
print("Local:", ARQUIVO_SAIDA)
